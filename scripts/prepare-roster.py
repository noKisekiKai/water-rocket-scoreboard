"""Read the teacher's source workbook; write a LOCAL-ONLY roster outside Git."""
import json, re, sys
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

source, destination = map(Path, sys.argv[1:3])
workbook = load_workbook(source, read_only=True, data_only=True)
groups, schools = {}, {}
for sheet in workbook:
    for row in sheet.iter_rows(values_only=True):
        for column, value in enumerate(row):
            if not isinstance(value, str) or not value.strip():
                continue
            value = value.strip()
            match = re.fullmatch(r'(\d+)\s*(\S.*)', value)
            if not match:
                if '學校' in value or '小學' in value:
                    schools[column] = value
                continue
            number, name = int(match[1]), match[2].strip()
            school = schools.get(column)
            if not school:
                raise ValueError(f'Missing school for group {number}')
            group = groups.setdefault(number, {'id': f'group-{number}', 'number': number, 'school': school, 'students': []})
            if group['school'] != school:
                raise ValueError(f'School mismatch for group {number}')
            if name in group['students']:
                raise ValueError(f'Duplicate name in group {number}')
            group['students'].append(name)
result = {'schema': 'water-rocket-roster-v1', 'groups': [groups[n] for n in sorted(groups)]}
destination.parent.mkdir(parents=True, exist_ok=True)
destination.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'groups': len(groups), 'students': sum(len(g['students']) for g in groups.values()), 'schools': len(set(g['school'] for g in groups.values())), 'groupSizes': dict(Counter(len(g['students']) for g in groups.values()))}, ensure_ascii=False))
